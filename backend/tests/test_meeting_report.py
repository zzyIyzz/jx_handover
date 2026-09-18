"""月度例会材料、全年定期工作一览与定期工作“记忆”的回归测试。"""
from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.db import Base
from app.models import (
    HandoverBatch,
    HandoverGeneralItem,
    HandoverItem,
    HandoverStationMeta,
    MonthlyPlanItem,
    Station,
    WorkItem,
)
from app.services import meeting_report


class MeetingReportTests(unittest.TestCase):
    def setUp(self) -> None:
        engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(engine)
        self.db = sessionmaker(bind=engine)()
        self.station = Station(code="TEST", name="测试场站", aliases_json='["测试"]')
        self.db.add(self.station)
        self.db.commit()

    def tearDown(self) -> None:
        self.db.close()

    def _batch_with_item(self, status: str, priority: str, section: str,
                         title: str = "检修事项"):
        batch = HandoverBatch(start_date="2026-09-01", end_date="2026-09-10",
                              handover_date="2026-09-10", status="published")
        self.db.add(batch)
        self.db.flush()
        meta = HandoverStationMeta(batch_id=batch.id, station_id=self.station.id,
                                   duty_leader="甲")
        self.db.add(meta)
        self.db.flush()
        work = WorkItem(station_id=self.station.id, canonical_title=title,
                        status=status, priority=priority,
                        first_seen_date="2026-09-01", last_seen_date="2026-09-10")
        self.db.add(work)
        self.db.flush()
        item = HandoverItem(batch_id=batch.id, station_meta_id=meta.id,
                            work_item_id=work.id, title_snapshot=title,
                            status=status, priority=priority, section=section,
                            completed_by="甲" if status == "completed" else "",
                            next_owner="乙")
        self.db.add(item)
        self.db.commit()
        return batch, meta, item

    def _yearly_plan(self, library_id: str, plan_month: str,
                     status: str | None = None, note: str = ""):
        plan = MonthlyPlanItem(plan_month=plan_month, scope_type="station",
                               station_id=self.station.id,
                               title=f"定期项{library_id}", category="yearly",
                               library_id=library_id)
        self.db.add(plan)
        self.db.flush()
        if status is not None:
            batch = HandoverBatch(start_date=f"{plan_month}-01",
                                  end_date=f"{plan_month}-10",
                                  handover_date=f"{plan_month}-10",
                                  status="published")
            self.db.add(batch)
            self.db.flush()
            meta = HandoverStationMeta(batch_id=batch.id,
                                       station_id=self.station.id,
                                       duty_leader="甲")
            self.db.add(meta)
            self.db.flush()
            self.db.add(HandoverGeneralItem(
                batch_id=batch.id, station_meta_id=meta.id,
                monthly_plan_item_id=plan.id, status=status,
                owner="周智源", note=note))
        self.db.commit()
        return plan

    def test_monthly_meeting_groups_by_status(self):
        self._batch_with_item("completed", "urgent", "important", "已完成重点")
        self._batch_with_item("in_progress", "normal", "handover", "进行中交接")
        self._batch_with_item("pending", "important", "handover", "未开展交接")

        data = meeting_report.monthly_meeting(self.db, "2026-09")
        self.assertEqual(data["month_label"], "2026 年 9 月")
        self.assertEqual(data["summary"]["completed"], 1)
        self.assertEqual(data["summary"]["in_progress"], 1)
        # pending 交接 + 当月应开展的年度定期工作（未记录）
        self.assertGreaterEqual(data["summary"]["pending"], 2)

        completed = next(g for g in data["groups"] if g["key"] == "completed")
        self.assertEqual(completed["rows"][0]["title"], "已完成重点")
        self.assertEqual(completed["rows"][0]["priority_label"], "紧急")
        self.assertEqual(completed["rows"][0]["owner"], "甲")

        pending = next(g for g in data["groups"] if g["key"] == "pending")
        titles = [r["title"] for r in pending["rows"]]
        self.assertIn("未开展交接", titles)
        # 9 月应开展的年度定期工作（如安全月、质量月）即使没记录也要出现
        self.assertIn("质量月", titles)

    def test_monthly_meeting_outside_month_is_empty(self):
        self._batch_with_item("completed", "normal", "important")
        data = meeting_report.monthly_meeting(self.db, "2026-10")
        self.assertEqual(data["batch_count"], 0)
        self.assertEqual(data["summary"]["completed"], 0)

    def test_monthly_meeting_rejects_bad_month(self):
        from fastapi import HTTPException
        with self.assertRaises(HTTPException):
            meeting_report.monthly_meeting(self.db, "2026/09")

    def test_export_xlsx(self):
        self._batch_with_item("completed", "urgent", "important")
        with tempfile.TemporaryDirectory() as tmp:
            from unittest import mock
            from app import config
            with mock.patch.object(config, "GENERATED_DIR", Path(tmp)):
                path = meeting_report.export_monthly_meeting_xlsx(self.db, "2026-09")
            self.assertTrue(path.exists())
            from openpyxl import load_workbook
            wb = load_workbook(path)
            ws = wb.active
            self.assertIn("月度例会工作材料", ws.cell(row=1, column=1).value)

    def test_yearly_overview_with_memory(self):
        # 2025 年留下历史记录（记忆），2026 年有一条已完成
        self._yearly_plan("y1", "2025-01", status="completed", note="去年已签")
        self._yearly_plan("y2", "2026-01", status="completed", note="今年已完成")
        self._yearly_plan("y3", "2026-01", status="pending")

        data = meeting_report.yearly_overview(self.db, 2026)
        self.assertEqual(data["summary"]["total"], 33)
        by_id = {it["library_id"]: it for it in data["items"]}

        # y1 今年没有记录 → 未开展，但能看到去年的记忆
        self.assertEqual(by_id["y1"]["state"], "pending")
        self.assertIsNotNone(by_id["y1"]["memory"])
        self.assertEqual(by_id["y1"]["memory"]["note"], "去年已签")

        # y2 今年已完成
        self.assertEqual(by_id["y2"]["state"], "completed")
        self.assertEqual(by_id["y2"]["latest"]["note"], "今年已完成")

        # y3 今年有记录但未完成 → 正在开展
        self.assertEqual(by_id["y3"]["state"], "in_progress")

    def test_periodic_memory_endpoint(self):
        self._yearly_plan("y1", "2025-01", status="completed", note="历史备注")
        result = meeting_report.periodic_memory(self.db, ["y1", "y99", "bad"])
        self.assertIn("y1", result["memory"])
        self.assertEqual(result["memory"]["y1"]["note"], "历史备注")
        self.assertNotIn("bad", result["memory"])


if __name__ == "__main__":
    unittest.main()
