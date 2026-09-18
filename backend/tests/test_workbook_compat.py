"""Regression tests for malformed styles produced by XLSX exporters."""
from __future__ import annotations

from io import BytesIO
from pathlib import Path
import re
import sys
import unittest
from unittest import mock
from types import SimpleNamespace
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook


PROJECT_ROOT = Path(__file__).resolve().parents[2]
BACKEND_ROOT = PROJECT_ROOT / "backend"
sys.path.insert(0, str(BACKEND_ROOT))

from app.services.importer import workbook as workbook_compat
from app.services.importer import sections


def _write_malformed_workbook(path: Path) -> None:
    normal = BytesIO()
    book = Workbook()
    book.active["A1"] = "兼容性测试"
    book.save(normal)
    normal.seek(0)

    with ZipFile(normal, "r") as source, ZipFile(
        path, "w", compression=ZIP_DEFLATED
    ) as target:
        for entry in source.infolist():
            payload = source.read(entry.filename)
            if entry.filename == "xl/styles.xml":
                payload, replacements = re.subn(
                    rb"<fill><patternFill/></fill>", b"<fill/>", payload, count=1
                )
                if replacements != 1:
                    raise AssertionError("test fixture could not locate a fill")
            target.writestr(entry, payload)


class WorkbookCompatibilityTest(unittest.TestCase):
    def test_malformed_empty_fill_is_repaired_in_memory(self) -> None:
        from tempfile import TemporaryDirectory

        with TemporaryDirectory() as tmp:
            source = Path(tmp) / "malformed.xlsx"
            _write_malformed_workbook(source)
            original_bytes = source.read_bytes()

            book = workbook_compat.load_workbook_compat(
                source, data_only=True, read_only=True
            )
            try:
                self.assertEqual(book.active["A1"].value, "兼容性测试")
            finally:
                book.close()

            self.assertEqual(source.read_bytes(), original_bytes)

    def test_repair_keeps_fill_count_and_positions(self) -> None:
        styles = (
            b'<styleSheet xmlns="http://schemas.openxmlformats.org/'
            b'spreadsheetml/2006/main"><fills count="3">'
            b'<fill><patternFill patternType="none"/></fill><fill/>'
            b'<fill><patternFill patternType="solid"/></fill>'
            b'</fills></styleSheet>'
        )

        repaired, count = workbook_compat._repair_styles_xml(styles)

        self.assertEqual(count, 1)
        self.assertIn(b'<fills count="3">', repaired)
        self.assertEqual(repaired.count(b"<fill>"), 3)
        self.assertIn(
            b'<fill><patternFill patternType="none"/></fill>', repaired
        )

    def test_valid_workbook_does_not_use_repair_path(self) -> None:
        from tempfile import TemporaryDirectory

        with TemporaryDirectory() as tmp:
            source = Path(tmp) / "valid.xlsx"
            book = Workbook()
            book.active["A1"] = "正常文件"
            book.save(source)

            self.assertIsNone(workbook_compat._repaired_workbook_stream(source))
            with mock.patch.object(
                workbook_compat,
                "_repaired_workbook_stream",
                return_value=None,
            ) as repair_builder:
                loaded = workbook_compat.load_workbook_compat(source)
            try:
                self.assertEqual(loaded.active["A1"].value, "正常文件")
            finally:
                loaded.close()
            repair_builder.assert_called_once_with(source)


class MonthlyShiftWorkbookTest(unittest.TestCase):
    def test_two_row_monthly_layout_is_detected_and_parsed(self) -> None:
        book = Workbook()
        sheet = book.active
        sheet.title = "2月工作"
        sheet.append(["日期", "场站", "早班会", None, None, None, None,
                      "车辆", None, None, "晚班会"])
        sheet.append([None, None, "工作内容", "安全风险辨识", "特殊事项安排",
                      "值班负责人", "工作责任人", "车辆", "司机",
                      "驾驶风险告知", "工作完成情况"])
        sheet.append(["2026-02-03", "测试光伏电站", "逆变器巡检", "触电",
                      "携带验电器", "值班甲", "责任乙", None, None, None,
                      "进行中"])
        # Date and station are merged-style blank cells and must inherit the
        # preceding values, matching the real exporter layout.
        sheet.append([None, None, "箱变测温", "高温", None, "值班甲", "责任丙",
                      None, None, None, "已完成"])
        sheet.append([None, "其他电站", "不应导入", None, None, None, None,
                      None, None, None, "已完成"])

        batch = SimpleNamespace(
            start_date="2026-02-01", end_date="2026-02-28"
        )
        station = SimpleNamespace(name="测试光伏电站", aliases_json="[]")

        self.assertEqual(sections._detect_adapter(book), "monthly_shift_log")
        rows, warnings = sections._parse_monthly_shift_log(book, batch, station)

        self.assertEqual([row["title_snapshot"] for row in rows], [
            "逆变器巡检", "箱变测温",
        ])
        self.assertEqual(rows[0]["start_date"], "2026-02-03")
        self.assertEqual(rows[0]["previous_owner"], "责任乙")
        self.assertIn("安全风险辨识：触电", rows[0]["summary"])
        self.assertEqual(rows[1]["status"], "completed")
        self.assertEqual(rows[1]["completed_by"], "责任丙")
        self.assertTrue(warnings)


if __name__ == "__main__":
    unittest.main()
